"""Word (.docx) and PowerPoint (.pptx) loaders."""

import re
from collections.abc import Iterator

from docx import Document as open_docx
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from langchain_core.documents import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.ingestion.errors import IngestionError
from app.ingestion.loaders.base import (
    FileLoader,
    HeadingStack,
    clean_cell,
    make_block,
    make_table_block,
    table_to_markdown,
)

_HEADING_STYLE = re.compile(r"^heading\s*(\d)", re.IGNORECASE)


class DocxLoader(FileLoader):
    """Paragraphs and tables in document order; Heading 1..9 / Title styles build the heading path."""

    def lazy_load(self) -> Iterator[Document]:
        """Walk the document body in order, yielding paragraphs and tables."""
        try:
            doc = open_docx(str(self.path))
        except Exception as exc:
            raise IngestionError(f"Cannot open DOCX: {exc}") from exc

        stack = HeadingStack()
        for element in doc.element.body.iterchildren():
            tag = element.tag.rsplit("}", 1)[-1]
            if tag == "p":
                paragraph = Paragraph(element, doc)
                text = paragraph.text.strip()
                if not text:
                    continue
                style = (paragraph.style.name if paragraph.style is not None else "") or ""
                if style.lower() == "title":
                    stack.push(1, text)
                elif m := _HEADING_STYLE.match(style):
                    stack.push(int(m.group(1)) + 1, text)
                else:
                    prefix = "- " if style.lower().startswith("list") else ""
                    yield make_block(prefix + text, heading_path=stack.path)
            elif tag == "tbl":
                rows = [[clean_cell(cell.text) for cell in row.cells] for row in DocxTable(element, doc).rows]
                rows = [r for r in rows if any(r)]
                if len(rows) >= 2:
                    yield make_table_block(rows[0], rows[1:], heading_path=stack.path)
                elif rows:
                    yield make_block(" | ".join(c for c in rows[0] if c), heading_path=stack.path)


class PptxLoader(FileLoader):
    """One block per slide: title, body text (with bullet levels), tables and speaker notes."""

    def lazy_load(self) -> Iterator[Document]:
        """Yield one slide block per slide with its title, bullets, tables and notes."""
        try:
            prs = Presentation(str(self.path))
        except Exception as exc:
            raise IngestionError(f"Cannot open PPTX: {exc}") from exc
        self.page_count = len(prs.slides)

        for number, slide in enumerate(prs.slides, start=1):
            title_shape = slide.shapes.title
            title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
            parts: list[str] = []
            for shape in self._walk(slide.shapes):
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = " ".join(para.text.split())
                        if text:
                            parts.append("  " * para.level + "- " + text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    rows = [[clean_cell(c.text) for c in row.cells] for row in shape.table.rows]
                    if rows:
                        parts.append(table_to_markdown(rows[0], rows[1:]))
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"Speaker notes: {notes}")
            body = "\n".join(parts)
            if not (title or body):
                continue
            content = f"# {title}\n{body}".strip() if title else body
            yield make_block(content, "slide", heading_path=[title or f"Slide {number}"], page=number)

    def _walk(self, shapes):
        """All shapes of a slide, looking inside grouped shapes."""
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from self._walk(shape.shapes)
            else:
                yield shape
