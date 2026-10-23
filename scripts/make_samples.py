"""Create sample HRD documents (DOCX, XLSX, PPTX, PNG) in ./samples for trying the test UI.

Usage:  python scripts/make_samples.py
"""

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

OUT = Path(__file__).resolve().parents[1] / "samples"


def allowance_policy() -> None:
    """Word policy with headings, paragraphs and a table."""
    doc = Document()
    doc.add_heading("HRD Allowance Policy 2026", level=1)
    doc.add_paragraph("This policy describes the monthly allowances paid to HRD students.")
    doc.add_heading("Meal Allowance", level=2)
    doc.add_paragraph("Every active student receives a meal allowance of 60 USD per month.")
    doc.add_heading("IT Score Allowance", level=2)
    doc.add_paragraph("The IT score allowance depends on the monthly IT score:")
    table = doc.add_table(rows=4, cols=2)
    for r, (score, amount) in enumerate([("IT score", "Allowance"), ("90 and above", "30 USD"), ("80 - 89", "20 USD"), ("below 80", "0 USD")]):
        table.cell(r, 0).text, table.cell(r, 1).text = score, amount
    doc.add_heading("Korean Language Allowance", level=2)
    doc.add_paragraph("Students who pass the TOPIK level 2 test receive a Korean language allowance of 25 USD per month.")
    doc.add_heading("Extra-class Allowance", level=2)
    doc.add_paragraph("Extra-class hours are paid at 2 USD per hour, recorded by the instructor on standby days.")
    doc.save(OUT / "HRD Allowance Policy 2026.docx")


def attendance_sheet() -> None:
    """Excel attendance sheet with a title row above the header."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Gen 12 October"
    ws.append(["Generation 12 - October 2026 attendance"])
    ws.append(["Student", "Class", "Present days", "Absent days", "Late days", "Attendance rate"])
    students = [
        ("Sok Dara", "PP", 20, 0, 1), ("Chan Sreypov", "PP", 18, 2, 0), ("Kim Vibol", "SR", 14, 6, 3),
        ("Heng Malis", "SR", 21, 0, 0), ("Noun Rotha", "PP", 12, 8, 2), ("Lim Sopheak", "SR", 19, 1, 4),
    ]
    for name, cls, present, absent, late in students:
        ws.append([name, cls, present, absent, late, f"{present / (present + absent):.0%}"])
    wb.save(OUT / "Attendance Gen 12 October.xlsx")


def lesson_slides() -> None:
    """PowerPoint lesson with speaker notes."""
    prs = Presentation()
    slides = [
        ("Spring Boot REST API", "Build REST endpoints with @RestController\nUse @GetMapping and @PostMapping", "Show a live demo"),
        ("Spring Data JPA", "Entities map to tables\nRepositories give CRUD methods for free", "Explain JpaRepository"),
        ("Homework", "Build a student CRUD API with validation\nDeadline: Friday 5 PM", "Remind students about the deadline"),
    ]
    for title, body, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(OUT / "Lesson 5 - Spring Boot.pptx")


def scanned_notice() -> None:
    """PNG 'scan' of a notice, to exercise OCR."""
    image = Image.new("RGB", (1400, 700), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 40)
    except OSError:
        font = ImageFont.load_default(size=40)
    lines = [
        "HRD CENTER NOTICE",
        "The final project presentation will be held",
        "on 12 December 2026 at 9:00 AM in Room 301.",
        "Each team presents for 20 minutes.",
        "Contact: hrd.office@example.org",
    ]
    for i, line in enumerate(lines):
        draw.text((80, 80 + i * 100), line, fill="black", font=font)
    image.save(OUT / "Notice - Final Presentation.png")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    allowance_policy()
    attendance_sheet()
    lesson_slides()
    scanned_notice()
    print(f"Samples written to {OUT}")
